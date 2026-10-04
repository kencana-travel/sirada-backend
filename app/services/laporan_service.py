"""
Pembuatan laporan analisis (PDF via reportlab, Excel via openpyxl).

Laporan disusun dulu sebagai struktur netral (meta, KPI, daftar bagian bertabel), lalu
dirender ke PDF atau XLSX. Tiap bagian dibungkus try/except: bila satu analisis gagal (mis.
data belum cukup), bagian itu ditulis "data belum tersedia" dan laporan tetap jadi.
"""
import io
import logging
from datetime import date, datetime, timedelta, timezone

from sqlalchemy import case, func, select
from sqlalchemy.orm import Session

from app.models.models import Rute, Transaksi

log = logging.getLogger(__name__)

JENIS_LAPORAN = ("ringkasan", "performa", "segmentasi", "forecasting")
FORMAT_LAPORAN = ("pdf", "xlsx")
JUDUL_JENIS = {
    "ringkasan": "Ringkasan Lengkap",
    "performa": "Performa Rute & Cabang",
    "segmentasi": "Segmentasi Pelanggan (RFM + K-Means)",
    "forecasting": "Forecasting Permintaan 30 Hari",
}
WIB = timezone(timedelta(hours=7))
PESAN_KOSONG = "Data belum tersedia"

# Tipe kolom untuk format angka: "teks" | "int" | "rp" | "persen" | "desimal"


def _bagian(judul, kolom, tipe, baris, catatan=None):
    return {"judul": judul, "kolom": kolom, "tipe": tipe, "baris": baris, "catatan": catatan,
            "error": None}


def _bagian_gagal(judul, e: Exception | str):
    pesan = str(e) if isinstance(e, str) else f"{PESAN_KOSONG} ({e.__class__.__name__}: {e})"
    log.warning("Bagian laporan '%s' gagal: %s", judul, pesan)
    return {"judul": judul, "kolom": [], "tipe": [], "baris": [], "catatan": None,
            "error": pesan[:300]}


# --------------------------------------------------------------------------- KPI
def kpi_transaksi(db: Session, mulai: date | None, selesai: date | None,
                  cabang: str | None) -> list[tuple[str, float, str]]:
    penumpang = Transaksi.jenis_transaksi == "Penumpang"
    q = select(
        func.count(Transaksi.id_transaksi),
        func.sum(case((penumpang, 1), else_=0)),
        func.sum(case((penumpang, Transaksi.jumlah_unit), else_=0)),
        func.sum(case((penumpang, 0), else_=Transaksi.jumlah_unit)),
        func.sum(Transaksi.total_bayar),
        func.sum(case((penumpang, Transaksi.total_bayar), else_=0)),
        func.sum(case((penumpang, 0), else_=Transaksi.total_bayar)),
        func.count(func.distinct(Transaksi.tanggal)),
    )
    if mulai:
        q = q.where(Transaksi.tanggal >= mulai)
    if selesai:
        q = q.where(Transaksi.tanggal <= selesai)
    if cabang:
        q = q.where(Transaksi.cabang_asal == cabang)
    n, n_pnp, pnp, kg, total, rp_pnp, rp_paket, hari = db.execute(q).one()
    hari = hari or 0
    total = float(total or 0)
    return [
        ("Total transaksi", n or 0, "int"),
        ("Transaksi penumpang", n_pnp or 0, "int"),
        ("Total penumpang (orang)", pnp or 0, "int"),
        ("Total berat paket (kg)", kg or 0, "int"),
        ("Total pendapatan", total, "rp"),
        ("Pendapatan penumpang", float(rp_pnp or 0), "rp"),
        ("Pendapatan paket", float(rp_paket or 0), "rp"),
        ("Hari beroperasi", hari, "int"),
        ("Rata-rata pendapatan per hari", total / hari if hari else 0, "rp"),
    ]


# --------------------------------------------------------------------------- bagian analisis
def _bagian_performa(db, mulai, selesai, cabang) -> list[dict]:
    from app.services import performa_service as ps
    hasil = []
    try:
        data = ps.analisis_performa_rute(db, cabang=cabang, tanggal_mulai=mulai,
                                         tanggal_selesai=selesai)
        if not data:
            raise ValueError("tidak ada transaksi penumpang pada periode ini")
        hasil.append(_bagian(
            "Performa per Rute",
            ["Rute", "Total Trip", "Penumpang", "Pendapatan", "Okupansi (%)"],
            ["teks", "int", "int", "rp", "persen"],
            [[d.get("rute"), d.get("total_trip"), d.get("total_penumpang"),
              d.get("total_pendapatan"), d.get("okupansi_persen")] for d in data],
            "Okupansi = rata-rata okupansi per perjalanan (penumpang / kapasitas armada x 100%)."))
    except Exception as e:  # noqa: BLE001
        hasil.append(_bagian_gagal("Performa per Rute", e))
    try:
        data = ps.analisis_performa_cabang(db, cabang=cabang, tanggal_mulai=mulai,
                                           tanggal_selesai=selesai)
        if not data:
            raise ValueError("tidak ada transaksi penumpang pada periode ini")
        hasil.append(_bagian(
            "Performa per Cabang",
            ["Cabang", "Total Trip", "Penumpang", "Pendapatan", "Okupansi (%)"],
            ["teks", "int", "int", "rp", "persen"],
            [[d.get("cabang"), d.get("total_trip"), d.get("total_penumpang"),
              d.get("total_pendapatan"), d.get("okupansi_persen")] for d in data]))
    except Exception as e:  # noqa: BLE001
        hasil.append(_bagian_gagal("Performa per Cabang", e))
    return hasil


def _bagian_segmentasi(db, cabang) -> list[dict]:
    from app.services import segmentasi_service as ss
    hasil = []
    try:
        data = ss.jalankan_rfm_kmeans(db, n_clusters=None, cabang=cabang)
        ringkasan = data.get("ringkasan_cluster") or []
        if not ringkasan:
            raise ValueError(data.get("pesan") or "belum ada pelanggan untuk disegmentasi")
        hasil.append(_bagian(
            "Ringkasan Cluster Pelanggan",
            ["Segmen", "Jumlah Pelanggan", "Rata Recency (hari)", "Rata Frequency",
             "Rata Monetary"],
            ["teks", "int", "desimal", "desimal", "rp"],
            [[c.get("label"), c.get("jumlah_pelanggan"), c.get("rata_recency_hari"),
              c.get("rata_frequency"), c.get("rata_monetary")] for c in ringkasan],
            "Segmentasi memakai seluruh riwayat transaksi penumpang (tidak dibatasi periode)."))
        ev = data.get("evaluasi") or {}
        valid = ev.get("valid")
        hasil.append(_bagian(
            "Evaluasi Clustering",
            ["Metrik", "Nilai"], ["teks", "teks"],
            [["Jumlah cluster (k) terpilih", _teks(ev.get("k_terpilih"))],
             ["Silhouette Score (makin mendekati 1 makin baik)", _teks(ev.get("silhouette"))],
             ["Davies-Bouldin Index (makin kecil makin baik)", _teks(ev.get("dbi"))],
             ["Status validasi", "-" if valid is None else ("Valid" if valid else "Belum valid")]]))
    except Exception as e:  # noqa: BLE001
        hasil.append(_bagian_gagal("Segmentasi Pelanggan", e))
    return hasil


def _forecast_satu(fs, db, nama_rute):
    try:
        return fs.jalankan_forecast(db, nama_rute, horizon_hari=30, pakai_kalender=True)
    except TypeError as e:
        # Kompatibel dengan versi service yang belum punya parameter pakai_kalender.
        if "pakai_kalender" not in str(e):
            raise
        return fs.jalankan_forecast(db, nama_rute, horizon_hari=30)


def _bagian_forecasting(db, cabang) -> list[dict]:
    from app.services import forecasting_service as fs
    try:
        q = select(Rute.nama_rute).order_by(Rute.nama_rute)
        if cabang:
            q = q.where(Rute.cabang_asal == cabang)
        rute_list = list(db.scalars(q))
        if not rute_list:
            raise ValueError("tidak ada rute")
        baris, gagal = [], []
        for nama in rute_list:
            try:
                f = _forecast_satu(fs, db, nama)
                valid = f.get("valid")
                baris.append([nama, f.get("model"), f.get("mape_persen"), f.get("mae"),
                              f.get("rmse"),
                              "-" if valid is None else ("Ya" if valid else "Tidak"),
                              f.get("prediksi_periode_berikutnya")])
            except Exception as e:  # noqa: BLE001
                gagal.append(nama)
                baris.append([nama, PESAN_KOSONG, None, None, None, "-", None])
                log.warning("Forecast rute %s gagal: %s", nama, e)
        catatan = ("Model dipilih otomatis per rute; prediksi = total penumpang 30 hari ke depan. "
                   "Forecasting memakai seluruh riwayat (tidak dibatasi periode).")
        if gagal:
            catatan += f" Rute tanpa hasil: {', '.join(gagal)}."
        return [_bagian("Forecasting Permintaan per Rute",
                        ["Rute", "Model", "MAPE (%)", "MAE", "RMSE", "Valid",
                         "Prediksi 30 Hari"],
                        ["teks", "teks", "persen", "desimal", "desimal", "teks", "int"],
                        baris, catatan)]
    except Exception as e:  # noqa: BLE001
        return [_bagian_gagal("Forecasting Permintaan", e)]


def susun_laporan(db: Session, jenis: str, mulai: date | None, selesai: date | None,
                  cabang: str | None, dibuat_oleh: str) -> dict:
    if jenis not in JENIS_LAPORAN:
        raise ValueError(f"Jenis laporan tidak dikenal: {jenis}")
    bagian = []
    if jenis in ("ringkasan", "performa"):
        bagian += _bagian_performa(db, mulai, selesai, cabang)
    if jenis in ("ringkasan", "segmentasi"):
        bagian += _bagian_segmentasi(db, cabang)
    if jenis in ("ringkasan", "forecasting"):
        bagian += _bagian_forecasting(db, cabang)

    try:
        kpi = kpi_transaksi(db, mulai, selesai, cabang)
    except Exception as e:  # noqa: BLE001
        log.warning("KPI laporan gagal: %s", e)
        kpi = []

    if mulai or selesai:
        periode = f"{_tgl(mulai) if mulai else 'awal data'} s.d. {_tgl(selesai) if selesai else 'akhir data'}"
    else:
        periode = "Seluruh data"
    return {
        "judul": "Laporan Analisis SIRADA Kencana",
        "subjudul": JUDUL_JENIS[jenis],
        "jenis": jenis,
        "meta": [
            ("Jenis laporan", JUDUL_JENIS[jenis]),
            ("Periode", periode),
            ("Cabang", cabang or "Semua cabang"),
            ("Dibuat pada", datetime.now(WIB).strftime("%d-%m-%Y %H:%M WIB")),
            ("Dibuat oleh", dibuat_oleh),
        ],
        "kpi": kpi,
        "bagian": bagian,
    }


# --------------------------------------------------------------------------- format angka
_BULAN = ["Jan", "Feb", "Mar", "Apr", "Mei", "Jun", "Jul", "Agu", "Sep", "Okt", "Nov", "Des"]


def _tgl(d: date) -> str:
    return f"{d.day} {_BULAN[d.month - 1]} {d.year}"


def _teks(v) -> str:
    if v is None:
        return "-"
    if isinstance(v, float):
        return f"{v:.4f}".rstrip("0").rstrip(".")
    return str(v)


def _ribuan(n: float, desimal: int = 0) -> str:
    s = f"{n:,.{desimal}f}"
    return s.replace(",", "_").replace(".", ",").replace("_", ".")


def format_nilai(v, tipe: str) -> str:
    if v is None or v == "":
        return "-"
    if tipe == "teks":
        return str(v)
    try:
        n = float(v)
    except (TypeError, ValueError):
        return str(v)
    if tipe == "int":
        return _ribuan(round(n))
    if tipe == "rp":
        return "Rp " + _ribuan(round(n))
    if tipe == "persen":
        return _ribuan(n, 1) + "%"
    return _ribuan(n, 2)


# --------------------------------------------------------------------------- PDF
def render_pdf(lap: dict) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_RIGHT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import (KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table,
                                    TableStyle)
    from xml.sax.saxutils import escape

    biru = colors.HexColor("#1e3a8a")
    abu = colors.HexColor("#f1f5f9")
    garis = colors.HexColor("#cbd5e1")
    ss = getSampleStyleSheet()
    s_judul = ParagraphStyle("judul", parent=ss["Title"], textColor=biru, fontSize=18,
                             spaceAfter=2, alignment=0)
    s_sub = ParagraphStyle("sub", parent=ss["Heading2"], textColor=colors.HexColor("#334155"),
                           fontSize=12, spaceBefore=0, spaceAfter=8)
    s_h = ParagraphStyle("h", parent=ss["Heading3"], textColor=biru, fontSize=11.5,
                         spaceBefore=12, spaceAfter=4)
    s_sel = ParagraphStyle("sel", parent=ss["BodyText"], fontSize=8.5, leading=10.5)
    s_sel_r = ParagraphStyle("selr", parent=s_sel, alignment=TA_RIGHT)
    s_head = ParagraphStyle("head", parent=s_sel, textColor=colors.white,
                            fontName="Helvetica-Bold")
    s_cat = ParagraphStyle("cat", parent=ss["BodyText"], fontSize=8,
                           textColor=colors.HexColor("#64748b"), spaceBefore=3)
    s_err = ParagraphStyle("err", parent=ss["BodyText"], fontSize=9,
                           textColor=colors.HexColor("#b45309"))

    lebar = A4[0] - 3.6 * cm
    buf = io.BytesIO()

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(colors.HexColor("#64748b"))
        canvas.drawString(1.8 * cm, 1.1 * cm, f"{lap['judul']} - {lap['subjudul']}")
        canvas.drawRightString(A4[0] - 1.8 * cm, 1.1 * cm, f"Halaman {doc.page}")
        canvas.restoreState()

    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=1.8 * cm, rightMargin=1.8 * cm,
                            topMargin=1.6 * cm, bottomMargin=1.8 * cm,
                            title=lap["judul"], author="SIRADA Kencana")
    el = [Paragraph(escape(lap["judul"]), s_judul), Paragraph(escape(lap["subjudul"]), s_sub)]

    meta = Table([[Paragraph(f"<b>{escape(k)}</b>", s_sel), Paragraph(escape(str(v)), s_sel)]
                  for k, v in lap["meta"]], colWidths=[4 * cm, lebar - 4 * cm])
    meta.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), abu),
                              ("BOX", (0, 0), (-1, -1), 0.5, garis),
                              ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                              ("TOPPADDING", (0, 0), (-1, -1), 3),
                              ("BOTTOMPADDING", (0, 0), (-1, -1), 3)]))
    el += [meta, Spacer(1, 6)]

    if lap["kpi"]:
        el.append(Paragraph("Indikator Utama Transaksi", s_h))
        sel = [[Paragraph(f"<font size=7.5 color='#64748b'>{escape(k)}</font><br/>"
                          f"<b>{escape(format_nilai(v, t))}</b>", s_sel)
                for k, v, t in lap["kpi"][i:i + 3]] for i in range(0, len(lap["kpi"]), 3)]
        for r in sel:
            r += [""] * (3 - len(r))
        t = Table(sel, colWidths=[lebar / 3] * 3)
        t.setStyle(TableStyle([("BOX", (0, 0), (-1, -1), 0.5, garis),
                               ("INNERGRID", (0, 0), (-1, -1), 0.5, garis),
                               ("TOPPADDING", (0, 0), (-1, -1), 5),
                               ("BOTTOMPADDING", (0, 0), (-1, -1), 5)]))
        el.append(t)

    for b in lap["bagian"]:
        judul = Paragraph(escape(b["judul"]), s_h)
        if b["error"]:
            el.append(KeepTogether([judul, Paragraph(escape(b["error"]), s_err)]))
            continue
        # Bobot lebar kolom: kolom teks pertama & kolom Rupiah lebih lebar.
        bobot = [1.8 if (i == 0 and t == "teks") else 1.4 if t == "rp" else 1.0
                 for i, t in enumerate(b["tipe"])]
        col_w = [lebar * w / sum(bobot) for w in bobot]
        data = [[Paragraph(escape(k), s_head) for k in b["kolom"]]]
        for r in b["baris"]:
            data.append([Paragraph(escape(format_nilai(v, t)), s_sel if t == "teks" else s_sel_r)
                         for v, t in zip(r, b["tipe"])])
        t = Table(data, colWidths=col_w, repeatRows=1)
        gaya = [("BACKGROUND", (0, 0), (-1, 0), biru),
                ("BOX", (0, 0), (-1, -1), 0.5, garis),
                ("LINEBELOW", (0, 0), (-1, -1), 0.25, garis),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE")]
        for i in range(2, len(data), 2):
            gaya.append(("BACKGROUND", (0, i), (-1, i), abu))
        t.setStyle(TableStyle(gaya))
        isi = [judul, t]
        if b["catatan"]:
            isi.append(Paragraph(escape(b["catatan"]), s_cat))
        # Tabel pendek dijaga satu halaman dengan judulnya; tabel panjang boleh terpotong
        # (header tabel diulang lewat repeatRows).
        if len(data) <= 25:
            el.append(KeepTogether(isi))
        else:
            el += isi
    doc.build(el, onFirstPage=footer, onLaterPages=footer)
    return buf.getvalue()


# --------------------------------------------------------------------------- XLSX
_FORMAT_XLSX = {"int": "#,##0", "rp": '"Rp" #,##0', "persen": '0.0"%"', "desimal": "#,##0.00"}


def render_xlsx(lap: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Ringkasan"
    tebal = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="1E3A8A")
    head_font = Font(bold=True, color="FFFFFF")

    ws["A1"] = lap["judul"]
    ws["A1"].font = Font(bold=True, size=15, color="1E3A8A")
    ws["A2"] = lap["subjudul"]
    ws["A2"].font = Font(bold=True, size=12)
    r = 4
    for k, v in lap["meta"]:
        ws.cell(r, 1, k).font = tebal
        ws.cell(r, 2, v)
        r += 1
    if lap["kpi"]:
        r += 1
        ws.cell(r, 1, "Indikator Utama Transaksi").font = Font(bold=True, size=12,
                                                               color="1E3A8A")
        r += 1
        for k, v, t in lap["kpi"]:
            ws.cell(r, 1, k)
            c = ws.cell(r, 2, v)
            c.number_format = _FORMAT_XLSX.get(t, "General")
            r += 1
    r += 1
    ws.cell(r, 1, "Isi laporan").font = Font(bold=True, size=12, color="1E3A8A")
    r += 1
    for b in lap["bagian"]:
        ws.cell(r, 1, b["judul"])
        ws.cell(r, 2, b["error"] or f"lihat sheet '{_nama_sheet(b['judul'])}'")
        r += 1
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 48

    dipakai = {"Ringkasan"}
    for b in lap["bagian"]:
        nama = _nama_sheet(b["judul"])
        while nama in dipakai:
            nama = nama[:28] + str(len(dipakai))
        dipakai.add(nama)
        s = wb.create_sheet(nama)
        s["A1"] = b["judul"]
        s["A1"].font = Font(bold=True, size=13, color="1E3A8A")
        if b["error"]:
            s["A3"] = b["error"]
            s.column_dimensions["A"].width = 80
            continue
        for j, k in enumerate(b["kolom"], start=1):
            c = s.cell(3, j, k)
            c.fill, c.font = head_fill, head_font
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for i, baris in enumerate(b["baris"], start=4):
            for j, (v, t) in enumerate(zip(baris, b["tipe"]), start=1):
                c = s.cell(i, j, v if v is not None else "-")
                if t in _FORMAT_XLSX and isinstance(v, (int, float)):
                    c.number_format = _FORMAT_XLSX[t]
        if b["catatan"]:
            s.cell(len(b["baris"]) + 5, 1, b["catatan"]).font = Font(italic=True,
                                                                    color="64748B")
        for j, k in enumerate(b["kolom"], start=1):
            panjang = max([len(str(k))] + [len(format_nilai(x[j - 1], b["tipe"][j - 1]))
                                           for x in b["baris"]])
            s.column_dimensions[get_column_letter(j)].width = min(45, max(12, panjang + 3))
        s.freeze_panes = "A4"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _nama_sheet(judul: str) -> str:
    bersih = "".join(ch for ch in judul if ch not in '[]:*?/\\')
    return bersih[:31]


def buat_file(db: Session, jenis: str, fmt: str, mulai: date | None, selesai: date | None,
              cabang: str | None, dibuat_oleh: str) -> tuple[bytes, str, str]:
    """Kembalikan (isi_file, nama_file, media_type)."""
    if fmt not in FORMAT_LAPORAN:
        raise ValueError("Format harus pdf atau xlsx")
    lap = susun_laporan(db, jenis, mulai, selesai, cabang, dibuat_oleh)
    stamp = datetime.now(WIB).strftime("%Y%m%d_%H%M")
    bag_cabang = f"_{cabang.lower().replace(' ', '-')}" if cabang else ""
    nama = f"laporan_{jenis}{bag_cabang}_{stamp}.{fmt}"
    if fmt == "pdf":
        return render_pdf(lap), nama, "application/pdf"
    return (render_xlsx(lap), nama,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
